"""Extraction du texte d'un CV — sans LLM.

Première étape du pipeline (file `heavy`) : sortir du PDF ou du DOCX un texte
brut exploitable. Volontairement déterministe et sans modèle, pour trois
raisons : c'est l'étape la plus fréquente (une par CV), la moins tolérante à
l'aléa, et la plus facile à rejouer en cas d'échec.

Un CV scanné ne rend aucun texte : plutôt que d'envoyer une page vide au
modèle — qui produirait un profil inventé —, la candidature est marquée
`needs_ocr` et sort du lot avec un motif explicite. L'OCR lui-même est hors
périmètre du mémoire (chap. 9, perspectives).
"""

import io
import re

import pdfplumber
from docx import Document

# En dessous de ce nombre de caractères utiles, il n'y a pas de CV : page
# scannée, PDF d'images, ou fichier quasi vide.
MIN_USEFUL_CHARS = 200

# Garde-fou sur la taille envoyée au modèle : un CV de 40 pages existe, mais
# au-delà on tronque plutôt que de faire exploser le coût et la latence.
MAX_TEXT_CHARS = 40000

_WHITESPACE = re.compile(r"[ \t ]+")
_BLANK_LINES = re.compile(r"\n{3,}")


class ExtractionError(Exception):
    """Fichier illisible : corrompu, tronqué, ou chiffré."""


class NeedsOCR(ExtractionError):
    """Fichier valide mais sans texte : CV scanné, PDF d'images.

    Distinct d'un fichier corrompu : le document est intact, c'est notre
    chaîne qui ne sait pas le lire. La candidature sort du lot avec ce motif
    plutôt que d'envoyer une page vide au modèle, qui inventerait un profil.
    """


def normalize(raw: str) -> str:
    """Espaces et lignes vides normalisés — le texte part dans un prompt."""
    text = raw.replace("\r\n", "\n").replace("\r", "\n")
    text = _WHITESPACE.sub(" ", text)
    text = "\n".join(line.strip() for line in text.split("\n"))
    return _BLANK_LINES.sub("\n\n", text).strip()


def _from_pdf(data: bytes) -> str:
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            pages = [page.extract_text() or "" for page in pdf.pages]
    except Exception as exc:  # pdfminer lève des exceptions très variées
        raise ExtractionError(f"PDF illisible : {type(exc).__name__}") from exc
    return "\n\n".join(pages)


def _from_docx(data: bytes) -> str:
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError(f"DOCX illisible : {type(exc).__name__}") from exc
    blocks = [p.text for p in document.paragraphs]
    # Beaucoup de CV rangent les dates et les postes dans des tableaux.
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    return "\n".join(blocks)


def extract_text(data: bytes, kind: str) -> str:
    """Texte normalisé d'un CV.

    Lève `ExtractionError` si le fichier est illisible, `NeedsOCR` s'il est
    valide mais ne contient pas de texte exploitable.
    """
    raw = _from_pdf(data) if kind == "pdf" else _from_docx(data)
    text = normalize(raw)
    if len(text) < MIN_USEFUL_CHARS:
        raise NeedsOCR(
            f"Texte insuffisant ({len(text)} caractères) : document scanné ou vide, "
            "reconnaissance optique nécessaire"
        )
    return text[:MAX_TEXT_CHARS]
